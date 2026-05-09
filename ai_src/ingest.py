# ingest.py - Extração de texto de documentos (PDF, Excel, etc.)
# Início da alteração

from pathlib import Path
from typing import Optional


def extrair_pdf(caminho: Path) -> str:
    """Extrai texto de PDF usando pypdf."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(caminho))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as e:
        raise RuntimeError(f"Erro ao ler PDF {caminho}: {e}")


def extrair_excel(caminho: Path) -> str:
    """Extrai texto de planilha Excel."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(caminho), read_only=True, data_only=True)
        partes = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                linha = " | ".join(str(c) if c is not None else "" for c in row)
                if linha.strip():
                    partes.append(linha)
        return "\n".join(partes)
    except Exception as e:
        raise RuntimeError(f"Erro ao ler Excel {caminho}: {e}")


def extrair_docx(caminho: Path) -> str:
    """Extrai texto de documento Word."""
    try:
        from docx import Document
        doc = Document(str(caminho))
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception as e:
        raise RuntimeError(f"Erro ao ler DOCX {caminho}: {e}")


def extrair_texto(caminho: Path) -> tuple[str, str]:
    """
    Extrai texto do documento conforme extensão.
    Retorna (texto, nome_arquivo).
    """
    suf = caminho.suffix.lower()
    if suf == ".pdf":
        texto = extrair_pdf(caminho)
    elif suf in (".xlsx", ".xls"):
        texto = extrair_excel(caminho)
    elif suf == ".docx":
        texto = extrair_docx(caminho)
    else:
        raise ValueError(f"Formato não suportado: {suf}")
    return texto, caminho.name

# Fim da alteração
